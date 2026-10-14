from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

REPO_ROOT = Path(__file__).resolve().parents[3]
SUBMISSION_DIR = REPO_ROOT / "docs" / "final" / "submission"
PANDOC = Path(r"C:\Users\lenovo\AppData\Local\Pandoc\pandoc.exe")
WKHTMLTOPDF = Path(r"C:\Program Files\wkhtmltopdf\bin\wkhtmltopdf.exe")
SOURCE_MD = SUBMISSION_DIR / "Food_Delivery_System_Final_Report.md"
DOCX_OUT = SUBMISSION_DIR / "Food_Delivery_System_Final_Report.docx"
PDF_OUT = SUBMISSION_DIR / "Food_Delivery_System_Final_Report.pdf"
CSS_PATH = SUBMISSION_DIR / "_report_styles.css"


def ensure_tools():
    if not PANDOC.exists():
        raise FileNotFoundError(f"Pandoc is not available at {PANDOC}.")
    if not WKHTMLTOPDF.exists():
        raise FileNotFoundError(f"wkhtmltopdf is not available at {WKHTMLTOPDF}.")
    if not SOURCE_MD.exists():
        raise FileNotFoundError(f"Markdown source not found at {SOURCE_MD}")


def write_css():
    css = """
    @page {
      size: A4;
      margin: 20mm 15mm 20mm 15mm;
      @bottom-center {
        content: counter(page);
        font-family: Arial, sans-serif;
        font-size: 10pt;
      }
    }
    html, body {
      margin: 0;
      padding: 0;
      background: #fff;
      color: #111;
      font-family: "Segoe UI", Arial, sans-serif;
      line-height: 1.5;
    }
    body {
      width: 100%;
      max-width: 210mm;
      margin: 0 auto;
    }
    h1, h2, h3, h4 {
      page-break-after: avoid;
      break-after: avoid;
      color: #1b1b1b;
      font-weight: 600;
    }
    h1 { font-size: 26pt; margin-top: 18pt; }
    h2 { font-size: 18pt; margin-top: 18pt; }
    h3 { font-size: 14pt; margin-top: 14pt; }
    p, li { font-size: 11.5pt; }
    table {
      width: 100%;
      border-collapse: collapse;
      margin: 12pt 0;
      page-break-inside: avoid;
      break-inside: avoid;
    }
    th, td {
      border: 1px solid #444;
      padding: 6pt 8pt;
      text-align: left;
      vertical-align: top;
      font-size: 10.5pt;
    }
    th {
      background: #f2f2f2;
      font-weight: 700;
    }
    img {
      max-width: 100%;
      height: auto;
      display: block;
      margin: 10pt auto;
      page-break-inside: avoid;
      break-inside: avoid;
    }
    figure {
      margin: 12pt 0;
      page-break-inside: avoid;
      break-inside: avoid;
    }
    figcaption {
      font-size: 10pt;
      text-align: center;
      color: #333;
    }
    pre {
      background: #f7f7f7;
      border: 1px solid #ddd;
      padding: 10pt;
      overflow-x: auto;
      font-size: 9.5pt;
      margin: 12pt 0;
    }
    code {
      font-family: "Consolas", "Courier New", monospace;
      font-size: 9.5pt;
    }
    """
    CSS_PATH.write_text(css, encoding="utf-8")


def run_pandoc_docx():
    cmd = [
        str(PANDOC),
        str(SOURCE_MD),
        "-o",
        str(DOCX_OUT),
        "--standalone",
        "--toc",
        "--from=gfm",
        "--metadata=title:Food Delivery System",
        "--metadata=subtitle:DCIT 204/308 Joint Semester Project",
        "--metadata=author:Team Firewall",
        "--metadata=date:2026-10-14",
    ]
    subprocess.run(cmd, check=True, cwd=str(SUBMISSION_DIR))
    add_docx_page_numbers()


def add_docx_page_numbers():
    """Add a centred dynamic PAGE field to every report section footer."""
    document = Document(DOCX_OUT)
    for section in document.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        paragraph._p.append(field)
    document.save(DOCX_OUT)


def run_pandoc_pdf():
    cmd = [
        str(PANDOC),
        str(SOURCE_MD),
        "-o",
        str(PDF_OUT),
        "--standalone",
        "--toc",
        "--pdf-engine",
        str(WKHTMLTOPDF),
        "--from=gfm",
        "--css",
        str(CSS_PATH),
        "--pdf-engine-opt=--footer-center",
        "--pdf-engine-opt=[page]",
        "--pdf-engine-opt=--footer-font-name",
        "--pdf-engine-opt=Arial",
        "--pdf-engine-opt=--footer-font-size",
        "--pdf-engine-opt=10",
        "--pdf-engine-opt=--footer-spacing",
        "--pdf-engine-opt=8",
        "--metadata=title:Food Delivery System",
        "--metadata=subtitle:DCIT 204/308 Joint Semester Project",
        "--metadata=author:Team Firewall",
        "--metadata=date:2026-10-14",
    ]
    subprocess.run(cmd, check=True, cwd=str(SUBMISSION_DIR))


def ensure_outputs_exist():
    if not DOCX_OUT.exists() or DOCX_OUT.stat().st_size < 5000:
        raise RuntimeError(f"DOCX was not generated at {DOCX_OUT}")
    if not PDF_OUT.exists() or PDF_OUT.stat().st_size < 1000:
        raise RuntimeError(f"PDF was not generated at {PDF_OUT}")


def main():
    ensure_tools()
    write_css()
    run_pandoc_docx()
    run_pandoc_pdf()
    if CSS_PATH.exists():
        CSS_PATH.unlink()
    ensure_outputs_exist()
    print(f"DOCX: {DOCX_OUT} ({DOCX_OUT.stat().st_size} bytes)")
    print(f"PDF: {PDF_OUT} ({PDF_OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
